"""export_v4.py — сводный датасет v4 = v3 + налоговая отчётность ФНС по НДФЛ (build_ndfl.py)."""
import os, shutil
from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

OUT = Path(os.environ.get('ALTAI_OUT', str(Path(__file__).resolve().parents[2] / 'data' / 'dataset_v3')))
V4 = OUT / 'v4'; V4.mkdir(exist_ok=True)

summ = pd.read_csv(OUT / 'ndfl_mo_summary.csv', sep=';', dtype={'oktmo': str})
mo = pd.read_csv(OUT / 'ndfl_mo_2022_2025.csv', sep=';', dtype={'oktmo': str})
reg = pd.read_csv(OUT / 'ndfl_regions_2016_2025.csv', sep=';')
v3 = pd.read_csv(OUT / 'mo_summary.csv', sep=';', dtype={'oktmo': str})

# МО_сводка v4 = v3 + НДФЛ
s4 = v3.merge(summ.drop(columns='mo'), on='oktmo', how='left')
s4.to_csv(V4 / 'mo_summary.csv', sep=';', index=False, encoding='utf-8-sig')

# панель: добавляем НДФЛ в длинном формате
P = pd.read_csv(OUT / 'mo_panel_long.csv', sep=';', dtype={'oktmo': str})
SRC = 'ФНС, формы 7-НДФЛ / 5-НДФЛ / «Если быть точным», ndfl v20260722 (по месту учёта налогового агента)'
spec = [('agents', 'НДФЛ: налоговых агентов (6-НДФЛ)', 'ед.'), ('income_7', 'НДФЛ: сумма дохода, выплаченная агентами МО', 'руб.'),
        ('tax_7', 'НДФЛ: сумма налога исчисленная', 'руб.'), ('records_7', 'НДФЛ: сведений о физлицах, получивших доход', 'ед.')]
rows = []
for col, ind, unit in spec:
    t = mo[['oktmo', 'mo', 'year', col]].dropna().rename(columns={col: 'value'})
    t = t.assign(group='Доходы (НДФЛ)', indicator=ind, unit=unit, source=SRC,
                 note='7-НДФЛ; по МО полный охват с 2022 г.; Чемальский 2024 исключён (сбой публикации)')
    rows.append(t)
P4 = pd.concat([P, pd.concat(rows)[P.columns]], ignore_index=True)
P4.to_csv(V4 / 'mo_panel_long.csv', sep=';', index=False, encoding='utf-8-sig')

# графики
plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 10})
COL = {'Республика Алтай': '#1f6f78', 'Республика Тыва': '#a86a12', 'Республика Хакасия': '#7a7f86'}
fig, ax = plt.subplots(1, 2, figsize=(12, 4.4))
for r, g in reg.groupby('region'):
    g = g.dropna(subset=['income_7_per_cap_month'])
    ax[0].plot(g.year, g.income_7_per_cap_month / 1000, marker='o', lw=2, color=COL[r], label=r.replace('Республика ', ''))
    g2 = reg[(reg.region == r)].dropna(subset=['earners_share'])
    ax[1].plot(g2.year, g2.earners_share * 100, marker='o', lw=2, color=COL[r], label=r.replace('Республика ', ''))
ax[0].set_title('Доходы, выплаченные налоговыми агентами,\nтыс. руб. на жителя в месяц (7-НДФЛ)')
ax[1].set_title('Получатели доходов от налоговых агентов,\n% населения (5-НДФЛ)')
for a in ax:
    a.grid(alpha=.3); a.spines[['top', 'right']].set_visible(False); a.legend(frameon=False)
fig.text(0.01, -0.02, 'Источник: ФНС (5-НДФЛ, 7-НДФЛ) / «Если быть точным»; население — Росстат на 1 января. Тыва 2016, 2018, 2021 — нет или выброс.', fontsize=8, color='#555')
fig.tight_layout(); fig.savefig(V4 / 'ndfl_regions.png', dpi=150, bbox_inches='tight'); plt.close(fig)

s = summ.sort_values('То же на жителя в месяц, 2025, руб.')
fig, ax = plt.subplots(figsize=(8, 5))
bars = ax.barh(s.mo, s['То же на жителя в месяц, 2025, руб.'] / 1000, color=['#1f6f78' if m in ('Майминский район', 'г. Горно-Алтайск') else '#9fb8bb' for m in s.mo])
for b, v in zip(bars, s['То же на жителя в месяц, 2025, руб.']):
    ax.text(b.get_width() + .4, b.get_y() + b.get_height() / 2, f'{v/1000:.1f}', va='center', fontsize=9)
ax.set_title('Доходы, выплаченные налоговыми агентами МО, 2025,\nтыс. руб. на жителя в месяц', loc='left')
ax.spines[['top', 'right']].set_visible(False); ax.grid(axis='x', alpha=.3)
fig.text(0.01, -0.03, 'По месту учёта налогового агента (работодателя). Источник: ФНС, 7-НДФЛ / «Если быть точным»; население на 1.01.2025 — Росстат.', fontsize=8, color='#555')
fig.tight_layout(); fig.savefig(V4 / 'ndfl_mo_2025.png', dpi=150, bbox_inches='tight'); plt.close(fig)

# workbook
shutil.copy(OUT / 'altai_dataset_v3.xlsx', V4 / 'altai_dataset_v4.xlsx')
wb = openpyxl.load_workbook(V4 / 'altai_dataset_v4.xlsx')
rd = wb['README']
rd['B2'] = 'Сводный рабочий датасет «Республика Алтай: пространственная диагностика», v4 (v3 + налоговая отчётность ФНС по НДФЛ: районы 2022–2025, сравнение с Тывой и Хакасией 2016–2025)'
rd.append(['Лист «НДФЛ_МО_сводка»', 'НДФЛ по 11 МО: доходы, выплаченные агентами МО, на жителя, рост 2022→2025, налог, агенты на 1000 жит., отношение ФОТ Росстата к доходам по 7-НДФЛ, получатели доходов 2024 (в т.ч. иностранцы, ИП)'])
rd.append(['Лист «НДФЛ_МО_ряды»', 'МО × год (2022–2025): агенты, доходы, налог, сведения о физлицах, население, ФОТ Росстата'])
rd.append(['Лист «НДФЛ_регионы»', 'Республика Алтай, Тыва, Хакасия, 2016–2025: доходы и налог на жителя, доля получателей доходов, средний доход на получателя, агенты на 1000 жит.'])
rd.append(['Важно про НДФЛ', 'Данные ФНС — по месту учёта налогового агента (работодателя), а не по месту жительства. Горно-Алтайск, Майма и Турочак завышены за счёт зарегистрированных там организаций; для районов это показатель «формальной экономики на территории», а не доходов жителей'])
hdr = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='1F6F78')


def put(name, df):
    if name in wb.sheetnames: del wb[name]
    ws = wb.create_sheet(name)
    ws.append(list(df.columns))
    for r in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in r])
    for c in ws[1]: c.font = hdr; c.fill = fill; c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[1].height = 60
    for i, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = 14 if i > 2 else 24
    ws.freeze_panes = 'C2'


RCOLS = {'region': 'Регион', 'year': 'Год', 'agents': 'Налоговых агентов', 'income_7': 'Доходы, выплаченные агентами (7-НДФЛ), руб.',
         'tax_7': 'НДФЛ исчисленный (7-НДФЛ), руб.', 'records_7': 'Сведений о физлицах (7-НДФЛ)', 'earners_5': 'Получателей доходов (5-НДФЛ)',
         'income_5': 'Сумма дохода (5-НДФЛ), руб.', 'pop_avg': 'Население на 1 января', 'income_7_per_cap_month': 'Доходы на жителя в месяц, руб.',
         'tax_7_per_cap': 'НДФЛ на жителя в год, руб.', 'earners_share': 'Доля получателей доходов в населении',
         'income_per_earner_month': 'Доход на получателя в месяц (5-НДФЛ), руб.', 'agents_per_1000': 'Агентов на 1000 жит.'}
MCOLS = {'oktmo': 'ОКТМО', 'mo': 'МО', 'year': 'Год', 'agents': 'Налоговых агентов', 'income_7': 'Доходы, выплаченные агентами МО, руб.',
         'tax_7': 'НДФЛ исчисленный, руб.', 'records_7': 'Сведений о физлицах', 'pop_jan': 'Население на 1 января',
         'fot_rosstat': 'ФОТ Росстата, тыс. руб.', 'income_7_per_cap_month': 'Доходы на жителя в месяц, руб.', 'tax_7_per_cap': 'НДФЛ на жителя, руб.',
         'agents_per_1000': 'Агентов на 1000 жит.', 'fot_to_ndfl': 'ФОТ Росстата / доходы по 7-НДФЛ'}
put('НДФЛ_МО_сводка', summ)
put('НДФЛ_МО_ряды', mo[list(MCOLS)].rename(columns=MCOLS).round(3))
put('НДФЛ_регионы', reg[list(RCOLS)].rename(columns=RCOLS).round(4))
if 'МО_сводка' in wb.sheetnames:
    put('МО_сводка', s4)
    wb.move_sheet('МО_сводка', offset=-(len(wb.sheetnames) - 3))
wb.save(V4 / 'altai_dataset_v4.xlsx')
for f in ['ndfl_mo_summary.csv', 'ndfl_mo_2022_2025.csv', 'ndfl_regions_2016_2025.csv', 'ndfl_build_log.txt']:
    shutil.copy(OUT / f, V4 / f)
print(wb.sheetnames)
